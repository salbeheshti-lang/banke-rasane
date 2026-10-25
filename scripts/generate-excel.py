#!/usr/bin/env python3
"""
Generate a professional Excel file from the BankResaneh talk shows data.
Includes all 560 shows with full details, formatting, and multiple sheets.
"""
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import re
import os
from datetime import datetime

# Helper: Remove illegal characters for Excel
def clean_for_excel(value):
    """Remove characters that Excel doesn't allow."""
    if value is None:
        return ''
    if not isinstance(value, str):
        value = str(value)
    # Remove illegal characters (0x00-0x08, 0x0B-0x0C, 0x0E-0x1F)
    illegal_chars = re.compile(r'[\x00-\x08\x0B\x0C\x0E-\x1F]')
    return illegal_chars.sub('', value)

# Load data
with open('/home/z/my-project/public/talk-shows-data.json', 'r', encoding='utf-8') as f:
    shows = json.load(f)

print(f'Loaded {len(shows)} shows')

# Create workbook
wb = openpyxl.Workbook()

# ===== Define styles =====
BRAND_GREEN = '029048'
BRAND_GREEN_DARK = '016035'
TEXT_DARK = '1A2E22'
TEXT_SOFT = '4A5D52'
TEXT_MUTED = '7A8A80'

header_font = Font(name='Vazirmatn', bold=True, color='FFFFFF', size=12)
header_fill = PatternFill(start_color=BRAND_GREEN, end_color=BRAND_GREEN, fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

title_font = Font(name='Vazirmatn', bold=True, color=BRAND_GREEN_DARK, size=18)
subtitle_font = Font(name='Vazirmatn', color=TEXT_SOFT, size=11)

data_font = Font(name='Vazirmatn', size=10, color=TEXT_DARK)
data_alignment = Alignment(vertical='center', wrap_text=True)
center_alignment = Alignment(horizontal='center', vertical='center')

thin_border = Border(
    left=Side(style='thin', color='D0D0D0'),
    right=Side(style='thin', color='D0D0D0'),
    top=Side(style='thin', color='D0D0D0'),
    bottom=Side(style='thin', color='D0D0D0')
)

alt_fill = PatternFill(start_color='F5F9F7', end_color='F5F9F7', fill_type='solid')

# ===== Sheet 1: خلاصه (Summary) =====
ws1 = wb.active
ws1.title = 'خلاصه'
ws1.sheet_view.rightToLeft = True

# Title
ws1['B2'] = 'بانک رسانه'
ws1['B2'].font = Font(name='Vazirmatn', bold=True, color=BRAND_GREEN_DARK, size=24)
ws1['B3'] = 'آرشیو تخصصی برنامه‌های تلویزیونی'
ws1['B3'].font = subtitle_font
ws1['B4'] = f'گزارش تولید شده در: {datetime.now().strftime("%Y/%m/%d")}'
ws1['B4'].font = Font(name='Vazirmatn', color=TEXT_MUTED, size=10)

# Stats
stats = [
    ('کل برنامه‌ها', len(shows), BRAND_GREEN),
    ('تعداد ژانرها', len(set(s['category'].split('،')[0] for s in shows)), '2563EB'),
    ('با توضیحات فارسی', sum(1 for s in shows if s.get('descriptionFa')), '7C3AED'),
    ('با لینک یوتیوب', sum(1 for s in shows if s.get('youtube')), 'DC2626'),
    ('با لینک IMDb', sum(1 for s in shows if s.get('imdb')), 'F59E0B'),
    ('با لینک ویکی‌پدیا', sum(1 for s in shows if s.get('wikipedia')), '0891B2'),
    ('با تصاویر', sum(1 for s in shows if s.get('images') and len(s['images']) > 0), 'DB2777'),
    ('با ویدیو', sum(1 for s in shows if s.get('videos') and len(s['videos']) > 0), 'EA580C'),
]

row = 6
for label, value, color in stats:
    ws1.cell(row=row, column=2, value=label).font = Font(name='Vazirmatn', bold=True, size=12, color=TEXT_DARK)
    ws1.cell(row=row, column=2).alignment = Alignment(horizontal='right', vertical='center')
    
    cell = ws1.cell(row=row, column=3, value=value)
    cell.font = Font(name='Vazirmatn', bold=True, size=14, color=color)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.fill = PatternFill(start_color=f'{color}15', end_color=f'{color}15', fill_type='solid')
    cell.border = thin_border
    
    row += 1

# Category distribution
row += 2
ws1.cell(row=row, column=2, value='توزیع دسته‌بندی‌ها').font = title_font
row += 1

categories = {}
for s in shows:
    cat = re.sub(r'^\d+\.\s*', '', s['category']).split('،')[0].strip()
    if cat not in categories:
        categories[cat] = 0
    categories[cat] += 1

# Header
for col, header in enumerate(['دسته‌بندی', 'تعداد', 'درصد'], 2):
    cell = ws1.cell(row=row, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

row += 1

for cat, count in sorted(categories.items(), key=lambda x: -x[1]):
    pct = count / len(shows) * 100
    
    ws1.cell(row=row, column=2, value=cat).font = data_font
    ws1.cell(row=row, column=2).alignment = Alignment(horizontal='right', vertical='center')
    ws1.cell(row=row, column=2).border = thin_border
    
    ws1.cell(row=row, column=3, value=count).font = Font(name='Vazirmatn', bold=True, size=11, color=BRAND_GREEN)
    ws1.cell(row=row, column=3).alignment = center_alignment
    ws1.cell(row=row, column=3).border = thin_border
    
    ws1.cell(row=row, column=4, value=f'{pct:.1f}%').font = data_font
    ws1.cell(row=row, column=4).alignment = center_alignment
    ws1.cell(row=row, column=4).border = thin_border
    
    row += 1

# Set column widths
ws1.column_dimensions['A'].width = 3
ws1.column_dimensions['B'].width = 40
ws1.column_dimensions['C'].width = 15
ws1.column_dimensions['D'].width = 15

# ===== Sheet 2: لیست کامل برنامه‌ها =====
ws2 = wb.create_sheet('لیست برنامه‌ها')
ws2.sheet_view.rightToLeft = True

# Headers
headers = [
    'شماره', 'نام برنامه', 'دسته‌بندی', 'منبع داده', 'اندازه استودیو',
    'تماشاگر', 'سبک دکور', 'جلوه‌ها', 'چیدمان', 'عناصر دکور',
    'مواد', 'رنگ‌آمیزی', 'نورپردازی', 'خلاصه دکور', 'توضیحات فارسی',
    'لینک یوتیوب', 'لینک IMDb', 'لینک ویکی‌پدیا', 'تعداد تصاویر', 'تعداد ویدیوها'
]

for col, header in enumerate(headers, 1):
    cell = ws2.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

# Data rows
for row_idx, show in enumerate(shows, 2):
    row_data = [
        show.get('num', ''),
        clean_for_excel(show.get('name', '')),
        clean_for_excel(show.get('category', '')),
        clean_for_excel(show.get('source', '')),
        clean_for_excel(show.get('size', '')),
        clean_for_excel(show.get('viewers', '')),
        clean_for_excel(show.get('style', '')),
        clean_for_excel(show.get('effects', '')),
        clean_for_excel(show.get('layout', '')),
        clean_for_excel(show.get('elements', '')),
        clean_for_excel(show.get('materials', '')),
        clean_for_excel(show.get('colors', '')),
        clean_for_excel(show.get('lighting', '')),
        clean_for_excel(show.get('summary', '')),
        clean_for_excel(show.get('descriptionFa', '')),
        clean_for_excel(show.get('youtube', '')),
        clean_for_excel(show.get('imdb', '')),
        clean_for_excel(show.get('wikipedia', '')),
        len(show.get('images', [])),
        len(show.get('videos', [])),
    ]
    
    for col, value in enumerate(row_data, 1):
        cell = ws2.cell(row=row_idx, column=col, value=value)
        cell.font = data_font
        cell.alignment = data_alignment
        cell.border = thin_border
        
        # Alternate row colors
        if row_idx % 2 == 0:
            cell.fill = alt_fill
        
        # Special formatting for name column (RTL)
        if col == 2:
            cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
            cell.font = Font(name='Vazirmatn', size=10, bold=True, color=TEXT_DARK)
        
        # Center alignment for number columns
        if col in [1, 19, 20]:
            cell.alignment = center_alignment

# Set column widths
col_widths = {
    'A': 8, 'B': 30, 'C': 35, 'D': 12, 'E': 18,
    'F': 15, 'G': 15, 'H': 15, 'I': 15, 'J': 25,
    'K': 20, 'L': 20, 'M': 15, 'N': 40, 'O': 50,
    'P': 30, 'Q': 30, 'R': 30, 'S': 12, 'T': 12,
}

for col, width in col_widths.items():
    ws2.column_dimensions[col].width = width

# Freeze top row
ws2.freeze_panes = 'A2'

# Auto-filter
ws2.auto_filter.ref = f'A1:T{len(shows)+1}'

# Set row height for header
ws2.row_dimensions[1].height = 40

# ===== Sheet 3: آمار دکور =====
ws3 = wb.create_sheet('آمار دکور')
ws3.sheet_view.rightToLeft = True

ws3['B2'] = 'تحلیل دکور و صحنه'
ws3['B2'].font = title_font
ws3['B3'] = 'آمار عناصر، مصالح، رنگ‌ها و نورپردازی'
ws3['B3'].font = subtitle_font

# Analyze decor elements
def count_items(field):
    counts = {}
    for s in shows:
        val = s.get(field, '')
        if not val:
            continue
        items = re.split(r'[,،]', val)
        for item in items:
            item = item.strip().lower()
            if item:
                counts[item] = counts.get(item, 0) + 1
    return sorted(counts.items(), key=lambda x: -x[1])[:15]

row = 5

# Elements
ws3.cell(row=row, column=2, value='عناصر دکور (پرتکرارترین)').font = Font(name='Vazirmatn', bold=True, size=14, color=BRAND_GREEN_DARK)
row += 1
for col, header in enumerate(['عنصر', 'تعداد'], 2):
    cell = ws3.cell(row=row, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border
row += 1
for item, count in count_items('elements'):
    ws3.cell(row=row, column=2, value=item).font = data_font
    ws3.cell(row=row, column=2).border = thin_border
    ws3.cell(row=row, column=3, value=count).font = Font(name='Vazirmatn', bold=True, color=BRAND_GREEN)
    ws3.cell(row=row, column=3).alignment = center_alignment
    ws3.cell(row=row, column=3).border = thin_border
    row += 1

row += 2

# Materials
ws3.cell(row=row, column=2, value='مصالح به‌کاررفته').font = Font(name='Vazirmatn', bold=True, size=14, color='2563EB')
row += 1
for col, header in enumerate(['ماده', 'تعداد'], 2):
    cell = ws3.cell(row=row, column=col, value=header)
    cell.font = header_font
    cell.fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    cell.alignment = header_alignment
    cell.border = thin_border
row += 1
for item, count in count_items('materials'):
    ws3.cell(row=row, column=2, value=item).font = data_font
    ws3.cell(row=row, column=2).border = thin_border
    ws3.cell(row=row, column=3, value=count).font = Font(name='Vazirmatn', bold=True, color='2563EB')
    ws3.cell(row=row, column=3).alignment = center_alignment
    ws3.cell(row=row, column=3).border = thin_border
    row += 1

row += 2

# Colors
ws3.cell(row=row, column=2, value='طیف رنگی').font = Font(name='Vazirmatn', bold=True, size=14, color='7C3AED')
row += 1
for col, header in enumerate(['رنگ', 'تعداد'], 2):
    cell = ws3.cell(row=row, column=col, value=header)
    cell.font = header_font
    cell.fill = PatternFill(start_color='7C3AED', end_color='7C3AED', fill_type='solid')
    cell.alignment = header_alignment
    cell.border = thin_border
row += 1
for item, count in count_items('colors'):
    ws3.cell(row=row, column=2, value=item).font = data_font
    ws3.cell(row=row, column=2).border = thin_border
    ws3.cell(row=row, column=3, value=count).font = Font(name='Vazirmatn', bold=True, color='7C3AED')
    ws3.cell(row=row, column=3).alignment = center_alignment
    ws3.cell(row=row, column=3).border = thin_border
    row += 1

row += 2

# Lighting
ws3.cell(row=row, column=2, value='نوع نورپردازی').font = Font(name='Vazirmatn', bold=True, size=14, color='EA580C')
row += 1
for col, header in enumerate(['نورپردازی', 'تعداد'], 2):
    cell = ws3.cell(row=row, column=col, value=header)
    cell.font = header_font
    cell.fill = PatternFill(start_color='EA580C', end_color='EA580C', fill_type='solid')
    cell.alignment = header_alignment
    cell.border = thin_border
row += 1
for item, count in count_items('lighting'):
    ws3.cell(row=row, column=2, value=item).font = data_font
    ws3.cell(row=row, column=2).border = thin_border
    ws3.cell(row=row, column=3, value=count).font = Font(name='Vazirmatn', bold=True, color='EA580C')
    ws3.cell(row=row, column=3).alignment = center_alignment
    ws3.cell(row=row, column=3).border = thin_border
    row += 1

ws3.column_dimensions['A'].width = 3
ws3.column_dimensions['B'].width = 35
ws3.column_dimensions['C'].width = 15

# ===== Sheet 4: نمایه الفبایی =====
ws4 = wb.create_sheet('نمایه الفبایی')
ws4.sheet_view.rightToLeft = True

# Headers
ws4_headers = ['شماره', 'نام برنامه', 'دسته‌بندی', 'سبک']
for col, header in enumerate(ws4_headers, 1):
    cell = ws4.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

# Sort shows alphabetically by name
sorted_shows = sorted(shows, key=lambda s: s.get('name', '').lower())

for row_idx, show in enumerate(sorted_shows, 2):
    ws4.cell(row=row_idx, column=1, value=show.get('num', '')).font = data_font
    ws4.cell(row=row_idx, column=1).alignment = center_alignment
    ws4.cell(row=row_idx, column=1).border = thin_border
    
    ws4.cell(row=row_idx, column=2, value=show.get('name', '')).font = Font(name='Vazirmatn', size=10, bold=True, color=TEXT_DARK)
    ws4.cell(row=row_idx, column=2).alignment = Alignment(horizontal='right', vertical='center')
    ws4.cell(row=row_idx, column=2).border = thin_border
    
    ws4.cell(row=row_idx, column=3, value=show.get('category', '')).font = data_font
    ws4.cell(row=row_idx, column=3).alignment = data_alignment
    ws4.cell(row=row_idx, column=3).border = thin_border
    
    ws4.cell(row=row_idx, column=4, value=show.get('style', '')).font = data_font
    ws4.cell(row=row_idx, column=4).alignment = center_alignment
    ws4.cell(row=row_idx, column=4).border = thin_border
    
    if row_idx % 2 == 0:
        for col in range(1, 5):
            ws4.cell(row=row_idx, column=col).fill = alt_fill

ws4.column_dimensions['A'].width = 10
ws4.column_dimensions['B'].width = 35
ws4.column_dimensions['C'].width = 40
ws4.column_dimensions['D'].width = 15
ws4.freeze_panes = 'A2'
ws4.auto_filter.ref = f'A1:D{len(sorted_shows)+1}'

# Save
output_path = '/home/z/my-project/download/banke-rasane-database.xlsx'
wb.save(output_path)

size = os.path.getsize(output_path)
print(f'\n✓ Excel file created: {output_path}')
print(f'  Size: {size:,} bytes ({size/1024:.1f} KB)')
print(f'  Sheets: {len(wb.sheetnames)}')
for name in wb.sheetnames:
    print(f'    - {name}')
